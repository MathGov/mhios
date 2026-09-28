"""Read-only packaging, navigation and reference identity checks; not a Core re-audit."""
from __future__ import annotations
import argparse
import hashlib
import html
import json
import re
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from docx import Document
import fitz
from .validate import load_json
ROOT=Path(__file__).resolve().parents[1]
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
S='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def tokens(x):return re.findall(r'\w+',html.unescape(x),re.UNICODE)
def audit(root=ROOT):
    checks=[]
    def add(name,ok,detail=None):checks.append({'check':name,'pass':bool(ok),'detail':detail})
    pins=load_json(root/'references/CORE_PIN_MANIFEST.json')
    add('Exactly fifteen pinned Core sources',len(pins['core_components'])==15)
    for entry in pins['core_components']:
        p=root/entry['path'];add('Core hash: '+entry['component'],p.is_file() and hashlib.sha256(p.read_bytes()).hexdigest()==entry['sha256'])
        if p.suffix=='.docx':
            d=Document(p);add('Core OOXML opens: '+entry['component'],len(d.paragraphs)>0)
    stem='MathGov_Human_Interface_and_Orchestration_Standard_v2.1'
    dp=root/'docs'/f'{stem}.docx';d=Document(dp)
    with ZipFile(dp) as z:
        add('MHIOS DOCX ZIP CRC',z.testzip() is None)
        r=E.fromstring(z.read('word/document.xml'));ns={'w':W}
        anchors=r.findall('.//w:bookmarkStart',ns);names=[n.get('{'+W+'}name') for n in anchors]
        links=r.findall('.//w:hyperlink',ns);internal=[n.get('{'+W+'}anchor') for n in links if n.get('{'+W+'}anchor')]
        add('Unique bookmark names',len(names)==len(set(names)))
        add('115 resolved contents targets',len(internal)==115 and all(a in names for a in internal))
        for part in ['word/header1.xml','word/footer1.xml']:
            text=' '.join(E.fromstring(z.read(part)).xpath('.//w:t/text()',namespaces=ns))
            add('No stale running identity: '+part,'v2.0' not in text and 'Final Complete 2026-09-04' not in text)
        add('No live PAGEREF',b'PAGEREF' not in z.read('word/document.xml'))
        nr=E.fromstring(z.read('word/numbering.xml'))
        restarts=nr.findall('.//w:startOverride',ns)
        add('Five independent list restarts at one',len(restarts)==5 and all(e.get('{'+W+'}val')=='1' for e in restarts))
        cap=[p for p in r.findall('w:body/w:p',ns) if ''.join(p.xpath('.//w:t/text()',namespaces=ns)).startswith('When terms such as intelligent,')]
        add('Capability examples use real italic runs, not literal asterisks',len(cap)==1 and len(cap[0].findall('.//w:i',ns))==2 and '*' not in ''.join(cap[0].xpath('.//w:t/text()',namespaces=ns)))
        app=E.fromstring(z.read('docProps/app.xml'))
        page_value=app.find('{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Pages')
        add('Stored MHIOS page count matches rendered PDF',page_value is not None and int(page_value.text)==len(fitz.open(root/'docs'/f'{stem}.pdf')))

    pieces=[];started=False
    for el in r.find('w:body',ns):
        text=' '.join(''.join(p.xpath('.//w:t/text()',namespaces=ns)) for p in ([el] if el.tag=='{'+W+'}p' else el.findall('.//w:p',ns)))
        if text.startswith('Edition, exact Core binding'):started=True
        if started:pieces.append(text)
    md=(root/'docs'/f'{stem}.md').read_text();md=md[md.index('Edition, exact Core binding'):]
    md=md.replace('<br>',' ').replace('\\|','|')
    add('Complete Markdown/DOCX body-and-table token parity',tokens(' '.join(pieces))==tokens(md))
    pdf=fitz.open(root/'docs'/f'{stem}.pdf');nav=[];layout=[]
    for i,p in enumerate(pdf):
        if not p.get_text().strip():layout.append((i+1,'blank'))
        for b in p.get_text('blocks'):
            if b[0]<-1 or b[1]<-1 or b[2]>p.rect.width+1 or b[3]>p.rect.height+1:layout.append((i+1,'out_of_page'))
        for link in p.get_links():
            if link['kind']==fitz.LINK_GOTO:
                label=''.join(tokens(p.get_textbox(link['from'])))
                target=pdf[link['page']]
                near=target.get_textbox(fitz.Rect(0,max(0,link['to'].y-8),target.rect.width,min(target.rect.height,link['to'].y+90)))
                good=bool(label) and label in ''.join(tokens(near))
                nav.append({'source_page':i+1,'target_page':link['page']+1,'label':p.get_textbox(link['from']),'pass':good})
    add('All rendered contents links land on matching headings',len(nav)==115 and all(x['pass'] for x in nav),{'count':len(nav),'failed':[x for x in nav if not x['pass']]})
    add('PDF has no blank pages or out-of-page text blocks',not layout,{'pages':len(pdf),'issues':layout})
    wp=root/'templates/MHIOS_v2.1_Implementation_Workbook.xlsx'
    with ZipFile(wp) as z:
        add('Preparation workbook ZIP CRC',z.testzip() is None)
        wr=E.fromstring(z.read('xl/workbook.xml'));sn={'s':S}
        add('Preparation workbook has fourteen retained sheets',len(wr.findall('.//s:sheet',sn))==14)
        fs=[];errors=[];cached=[]
        for f in z.namelist():
            if f.startswith('xl/worksheets/sheet') and f.endswith('.xml'):
                sr=E.fromstring(z.read(f))
                for cell in sr.findall('.//s:c',sn):
                    if cell.get('t')=='e':errors.append((f,cell.get('r')))
                    formula=cell.find('s:f',sn)
                    if formula is not None:
                        fs.append((f,cell.get('r'),formula.text));value=cell.find('s:v',sn);cached.append(value.text if value is not None else None)
        add('Twenty preparation formulas',len(fs)==20)
        add('No cached preparation formula errors',not errors)
        add('Untouched preparation rows remain UNASSESSED',all(v=='UNASSESSED' for v in cached),cached)
    return {'status':'PASS' if all(c['pass'] for c in checks) else 'FAIL','checks':checks,'contents_links':nav,'scope':'Artifact structure and supplied-source identity, not empirical or operational conformance.'}
def main():
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path);a=p.parse_args();r=audit()
    if a.output:a.output.write_text(json.dumps(r,indent=2)+'\n')
    print(json.dumps(r,indent=2));return 0 if r['status']=='PASS' else 1
if __name__=='__main__':raise SystemExit(main())
